"""Refresh the implementation appendix in the existing Word plan.
Run with the Codex bundled Python runtime. Render it with Word or the skill renderer before delivery.
"""
import json
from pathlib import Path
from docx import Document
from docx.shared import Pt

root = Path(__file__).resolve().parents[1]
target = root / 'docs' / 'PokemonPedia Improvement Plan.docx'
updates = json.loads((root / 'docs' / 'update-log.json').read_text(encoding='utf-8'))
verification = json.loads((root / 'docs' / 'verification.json').read_text(encoding='utf-8')) if (root / 'docs' / 'verification.json').exists() else {'unitTests': 19, 'browserTests': 17, 'auditVulnerabilities': 0, 'featuresThrough': 'F05'}
doc = Document(target)
for paragraph in list(doc.paragraphs):
    if paragraph.text.startswith('Implementation update:'):
        paragraph._p.getparent().remove(paragraph._p)
for border in list(doc.styles['Title'].element.xpath('.//w:pBdr')):
    border.getparent().remove(border)
for paragraph in doc.paragraphs:
    if paragraph.style.name == 'Title':
        for border in list(paragraph._p.xpath('.//w:pBdr')):
            border.getparent().remove(border)
body = doc._element.body
marker = next((p for p in doc.paragraphs if p.text == 'Implementation history'), None)
if marker is not None:
    deleting = False
    for child in list(body):
        if child is marker._p:
            deleting = True
        if deleting and not child.tag.endswith('sectPr'):
            body.remove(child)
heading = doc.add_paragraph('Implementation history', 'Heading 1')
heading.paragraph_format.page_break_before = True
doc.add_paragraph(f"Updated {verification.get('date', '11 October 2026')}. The preceding plan records the original audit and roadmap. This appendix records the work completed since that audit, with validation and limitations for every implementation batch. Current implementation status below supersedes historical statements that a feature is missing or a build fails.")
summary = doc.add_paragraph(f"Implementation update: The core improvements and roadmap features through {verification['featuresThrough']} are complete within the scopes recorded in the history. The cumulative suite has {verification['unitTests']} unit and API tests and {verification['browserTests']} browser scenarios, and the dependency audit reports {verification['auditVulnerabilities']} vulnerabilities. See Implementation history for each batch and its limits. Broader interface translation, public deployment, artwork-inclusive and field performance studies, and physical device checks remain. The findings below describe the original baseline.")
doc.paragraphs[1]._p.addnext(summary._p)
for entry in updates:
    heading = doc.add_paragraph(entry['step'], 'Heading 2')
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.page_break_before = bool(entry.get('startNewPage', False))
    status = doc.add_paragraph(f"{entry['date']} | {entry['status']}")
    status.paragraph_format.keep_with_next = True
    for line in entry['done']:
        doc.add_paragraph(line)
    p = doc.add_paragraph()
    p.add_run('Validation: ').bold = True
    p.add_run(entry['validation'])
    p = doc.add_paragraph()
    p.add_run('Limits and remaining work: ').bold = True
    p.add_run(entry['remaining'])
doc.add_paragraph('Continuing maintenance', 'Heading 2')
doc.add_paragraph('Start the next session with docs/next-session.md and docs/implementation-progress.md. Every subsequent implementation update must add or revise its step in docs/update-log.json, keep the handoff, progress record and README accurate, and refresh both this Word file and its matching PDF. Record what was done and what was verified; do not mark pending or untested work complete. AGENTS.md carries this project instruction for future updates.')
doc.core_properties.comments = 'Original audit preserved; cumulative implementation appendix added.'
doc.save(target)
print(f'Updated {target} with {len(updates)} implementation batches.')
